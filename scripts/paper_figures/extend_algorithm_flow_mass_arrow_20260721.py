#!/usr/bin/env python3
"""Compact the editable algorithm flow while preserving journal-size text."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches


ROOT = Path("/Users/saintyoung/Desktop/FOWT-CONTROLLER-main")
SOURCE = Path("/Users/saintyoung/Desktop/PPT/ppt素材.pptx")
OUT = ROOT / "outputs/paper_figures/20260721_algorithm_flow_arrow_extended"
FULL_PPTX = OUT / "ppt素材_算法框架图紧凑矩形框_字体统一_侧边反馈线优化.pptx"
SINGLE_PPTX = OUT / "算法框架图_紧凑矩形框_字体统一_侧边反馈线优化_可编辑.pptx"
TARGET_SLIDE_INDEX = 34


def find_shape(slide, name: str):
    matches = [shape for shape in slide.shapes if shape.name == name]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one shape named {name!r}, found {len(matches)}")
    return matches[0]


def remove_shape(slide, name: str) -> None:
    shape = find_shape(slide, name)
    shape._element.getparent().remove(shape._element)


def set_run_typeface(run) -> None:
    """Write explicit Office typefaces instead of inheriting theme fonts."""
    run_properties = run._r.get_or_add_rPr()
    has_chinese = any("\u3400" <= character <= "\u9fff" for character in run.text)
    primary_typeface = "Songti SC" if has_chinese else "Times New Roman"
    for tag, typeface in (
        # Chinese runs use Songti in every font slot because some renderers
        # incorrectly consult the Latin slot even for East-Asian glyphs.
        ("a:latin", primary_typeface),
        # Songti SC is the installed macOS name of the Chinese Song typeface;
        # using it explicitly prevents PowerPoint/LibreOffice theme fallback.
        ("a:ea", "Songti SC"),
        ("a:cs", primary_typeface),
    ):
        element = run_properties.find(qn(tag))
        if element is None:
            element = etree.Element(qn(tag))
            run_properties.append(element)
        element.set("typeface", typeface)


def normalize_text_fonts(slide) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                set_run_typeface(run)


def extend_arrow(presentation: Presentation) -> None:
    slide = presentation.slides[TARGET_SLIDE_INDEX]
    current_state = find_shape(slide, "node_current_state")
    wind_history = find_shape(slide, "node_wind_history")
    feedback_calc = find_shape(slide, "node_feedback_calc")
    forecast_calc = find_shape(slide, "node_forecast_calc")
    forecast_gate = find_shape(slide, "node_forecast_gate")
    fusion = find_shape(slide, "node_fusion")
    demand_gate = find_shape(slide, "node_demand_gate")
    previous_box = find_shape(slide, "node_previous_target")
    candidate_box = find_shape(slide, "node_candidate_target")
    target_box = find_shape(slide, "node_current_target")
    pump_update = find_shape(slide, "node_pump_state_update")
    updated_state = find_shape(slide, "node_updated_state")

    # Keep the original box sizes and type sizes.  Only tighten the vertical
    # rhythm so short process links do not consume unnecessary page height.
    feedback_calc.top = current_state.top + current_state.height + Inches(0.25)
    forecast_calc.top = wind_history.top + wind_history.height + Inches(0.25)
    forecast_gate.top = forecast_calc.top + forecast_calc.height + Inches(0.25)
    fusion.top = forecast_gate.top + forecast_gate.height + Inches(0.70)
    demand_gate.top = fusion.top + fusion.height + Inches(0.28)
    branch_top = demand_gate.top + demand_gate.height + Inches(0.22)
    previous_box.top = branch_top
    candidate_box.top = branch_top
    target_box.top = max(
        previous_box.top + previous_box.height,
        candidate_box.top + candidate_box.height,
    ) + Inches(0.22)
    pump_update.top = target_box.top + target_box.height + Inches(0.25)
    updated_state.top = pump_update.top + pump_update.height + Inches(0.25)

    def set_vertical(name: str, start: int, end: int) -> None:
        connector = find_shape(slide, name)
        connector.top = start
        connector.height = end - start
        if connector.height <= 0:
            raise RuntimeError(f"Connector {name!r} became non-positive")

    set_vertical(
        "a_state_feedback_1",
        current_state.top + current_state.height,
        feedback_calc.top,
    )
    set_vertical(
        "a_wind_forecast_1",
        wind_history.top + wind_history.height,
        forecast_calc.top,
    )
    set_vertical(
        "a_forecast_gate_1",
        forecast_calc.top + forecast_calc.height,
        forecast_gate.top,
    )

    fusion_mid = fusion.top + fusion.height // 2
    set_vertical(
        "a_feedback_fusion_1",
        feedback_calc.top + feedback_calc.height,
        fusion_mid,
    )
    find_shape(slide, "a_feedback_fusion_2").top = fusion_mid

    gate_mid = forecast_gate.top + forecast_gate.height // 2
    set_vertical("a_gate_no_2", gate_mid, fusion.top)
    find_shape(slide, "a_gate_no_1").top = gate_mid
    set_vertical(
        "a_gate_yes_1",
        forecast_gate.top + forecast_gate.height,
        fusion.top,
    )

    set_vertical(
        "a_fusion_threshold_1",
        fusion.top + fusion.height,
        demand_gate.top,
    )
    demand_mid = demand_gate.top + demand_gate.height // 2
    for horizontal_branch in (
        shape for shape in slide.shapes if shape.name == "a_threshold_no_1"
    ):
        horizontal_branch.top = demand_mid
    set_vertical("a_threshold_no_2", demand_mid, previous_box.top)
    # The source deck contains two shapes with this name; the zero-width
    # vertical one is the branch connector and the short horizontal one stays.
    threshold_yes = [
        shape
        for shape in slide.shapes
        if shape.name == "a_threshold_yes_1" and shape.width <= 1
    ][0]
    threshold_yes.top = demand_mid
    threshold_yes.height = candidate_box.top - demand_mid

    set_vertical(
        "a_previous_target_1",
        previous_box.top + previous_box.height,
        target_box.top,
    )
    set_vertical(
        "a_candidate_target_1",
        candidate_box.top + candidate_box.height,
        target_box.top,
    )
    set_vertical(
        "a_target_pump_1",
        target_box.top + target_box.height,
        pump_update.top,
    )
    set_vertical(
        "a_pump_state_1",
        pump_update.top + pump_update.height,
        updated_state.top,
    )

    # Reposition the decision labels inside the tightened branches.
    find_shape(slide, "label_gate_no").top = gate_mid + Inches(0.20)
    find_shape(slide, "label_zero_prediction").top = fusion.top - Inches(0.52)
    find_shape(slide, "label_gate_yes").top = forecast_gate.top + forecast_gate.height + Inches(0.04)
    find_shape(slide, "label_prediction_demand").top = fusion.top - Inches(0.34)
    find_shape(slide, "label_threshold_no").top = demand_mid + Inches(0.18)
    find_shape(slide, "label_threshold_yes").top = demand_mid + Inches(0.16)

    forecast_output = find_shape(slide, "label_forecast_output")
    forecast_output._element.getparent().remove(forecast_output._element)

    target_box.text_frame.paragraphs[0].runs[0].text = (
        "结合当前各舱压载质量确定本周期目标压载质量"
    )
    for name in (
        "node_current_mass",
        "a_mass_target_1",
        "a_mass_pump_1",
        "a_mass_pump_2",
    ):
        remove_shape(slide, name)

    updated_bottom = updated_state.top + updated_state.height
    updated_mid = updated_state.top + updated_state.height // 2
    current_state_mid = current_state.top + current_state.height // 2
    loop_state_bottom = find_shape(slide, "a_loop_state_1")
    loop_state_vertical = find_shape(slide, "a_loop_state_2")
    loop_state_target = find_shape(slide, "a_loop_state_3")
    loop_state_bottom.top = updated_mid
    loop_state_vertical.top = current_state_mid
    loop_state_vertical.height = updated_mid - current_state_mid
    loop_state_target.top = current_state_mid

    loop_bottom = find_shape(slide, "a_loop_mass_1")
    loop_vertical = find_shape(slide, "a_loop_mass_2")
    loop_target = find_shape(slide, "a_loop_mass_3")
    route_x = Inches(11.75)
    target_right = target_box.left + target_box.width
    target_mid = target_box.top + target_box.height // 2
    loop_bottom.top = updated_mid
    loop_bottom.width = route_x - loop_bottom.left
    loop_vertical.top = target_mid
    loop_vertical.height = updated_mid - target_mid
    loop_vertical.left = route_x
    loop_target.top = target_mid
    loop_target.left = target_right
    loop_target.width = route_x - target_right

    for shape in slide.shapes:
        preset_geometry = shape._element.spPr.prstGeom
        if preset_geometry is not None and preset_geometry.get("prst") == "roundRect":
            preset_geometry.set("prst", "rect")

    normalize_text_fonts(slide)


def retain_only_target_slide(presentation: Presentation) -> None:
    slide_id_list = presentation.slides._sldIdLst
    for index in reversed(range(len(presentation.slides))):
        if index == TARGET_SLIDE_INDEX:
            continue
        slide_id = slide_id_list[index]
        relationship_id = slide_id.rId
        presentation.part.drop_rel(relationship_id)
        slide_id_list.remove(slide_id)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(str(SOURCE))
    extend_arrow(presentation)
    presentation.save(str(FULL_PPTX))

    single = Presentation(str(FULL_PPTX))
    retain_only_target_slide(single)
    single.save(str(SINGLE_PPTX))
    print(FULL_PPTX)
    print(SINGLE_PPTX)


if __name__ == "__main__":
    main()
